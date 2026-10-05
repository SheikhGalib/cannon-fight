import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # blank layout

    # Color Palette: Deep Academic Navy / Dark Slate Theme
    BG_COLOR = RGBColor(15, 23, 42)        # Slate 900
    CARD_BG = RGBColor(30, 41, 59)        # Slate 800
    CARD_BORDER = RGBColor(51, 65, 85)    # Slate 700
    ACCENT_BLUE = RGBColor(56, 189, 248)  # Sky 400
    ACCENT_GOLD = RGBColor(251, 191, 36)  # Amber 400
    ACCENT_RED = RGBColor(248, 113, 113)  # Red 400
    ACCENT_GREEN = RGBColor(74, 222, 128) # Green 400
    TEXT_LIGHT = RGBColor(248, 250, 252)  # Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184)  # Slate 400
    WHITE = RGBColor(255, 255, 255)

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category="COMPUTER GRAPHICS PROJECT"):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_GOLD
        p_cat.font.name = "Arial"

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.7))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE
        p_title.font.name = "Arial"

        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.7), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = CARD_BORDER
        line.line.fill.background()

    def add_card(slide, left, top, width, height, bg_col=CARD_BG, border_col=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_col
        card.line.color.rgb = border_col
        card.line.width = Pt(1.2)
        return card

    # =========================================================================
    # SLIDE 1: Title & Introduction
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Accent decorative banner
    banner = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(0.12), Inches(3.8))
    banner.fill.solid()
    banner.fill.fore_color.rgb = ACCENT_GOLD
    banner.line.fill.background()

    t_box = s1.shapes.add_textbox(Inches(1.2), Inches(1.5), Inches(11.0), Inches(3.8))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "COMPUTER GRAPHICS LAB DEMONSTRATION"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GOLD

    p2 = tf1.add_paragraph()
    p2.text = "Medieval Fortress Siege Simulation"
    p2.font.size = Pt(36)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_before = Pt(8)

    p3 = tf1.add_paragraph()
    p3.text = "Hierarchical Kinematics, Gouraud & Blinn-Phong Shading, Soft Shadow Mapping,\nReal-Time Ray Tracing, and Ballistic Destruction in C++ / OpenGL 3.3 Core"
    p3.font.size = Pt(16)
    p3.font.color.rgb = ACCENT_BLUE
    p3.space_before = Pt(10)

    # Footer Card with project info
    add_card(s1, Inches(0.8), Inches(5.6), Inches(11.7), Inches(1.3))
    info_box = s1.shapes.add_textbox(Inches(1.0), Inches(5.7), Inches(11.3), Inches(1.1))
    tf_info = info_box.text_frame
    p_i1 = tf_info.paragraphs[0]
    p_i1.text = "• Core Deliverables: Transformations, Dual Shading (Gouraud & Phong), Lighting, Motion/Animation, Interaction, Ray Tracing"
    p_i1.font.size = Pt(12)
    p_i1.font.bold = True
    p_i1.font.color.rgb = TEXT_LIGHT

    p_i2 = tf_info.add_paragraph()
    p_i2.text = "• Framework: Pure C++17, OpenGL 3.3 Core, GLSL, GLFW 3.5, GLAD, GLM | Toolchain: WinLibs MinGW-w64 (GCC 16)"
    p_i2.font.size = Pt(11)
    p_i2.font.color.rgb = TEXT_MUTED
    p_i2.space_before = Pt(4)

    # =========================================================================
    # SLIDE 2: Proposal Scope vs. Final Implementation
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Scope Fulfillment: Proposal vs. Final Implementation")

    # Left Card: Proposed
    add_card(s2, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.3))
    box_p = s2.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(5.1))
    tf_p = box_p.text_frame
    tf_p.word_wrap = True
    p = tf_p.paragraphs[0]
    p.text = "📋 ORIGINAL PROPOSAL SCOPE"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GOLD

    items_proposed = [
        "Language & API: Python 3 / PyOpenGL Compatibility Profile",
        "Cannon Model: Single static/translating field cannon",
        "Kinematics: Wheel roll linked to distance, barrel elevation",
        "Shading Model: Basic Phong shading model with directional light",
        "Materials: 2 basic materials (matte wood carriage, shiny barrel)",
        "Projectile: Optional simple parabolic cannonball arc",
        "Target Scene: Flat ground plane with simple static target",
        "User Controls: Fixed viewpoint, basic keyboard input"
    ]
    for item in items_proposed:
        pi = tf_p.add_paragraph()
        pi.text = "• " + item
        pi.font.size = Pt(11)
        pi.font.color.rgb = TEXT_LIGHT
        pi.space_before = Pt(6)

    # Right Card: Delivered
    add_card(s2, Inches(6.9), Inches(1.7), Inches(5.6), Inches(5.3))
    box_d = s2.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.2), Inches(5.1))
    tf_d = box_d.text_frame
    tf_d.word_wrap = True
    p = tf_d.paragraphs[0]
    p.text = "🚀 DELIVERED FINAL IMPLEMENTATION"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    items_delivered = [
        "C++17 & OpenGL 3.3 Core Profile: Zero legacy GL; pure modern shaders",
        "Artillery Battery: 3 independently articulated cannons (drive, yaw, elevate)",
        "Dual Shading Models: Runtime toggle between Gouraud & Blinn-Phong ('G')",
        "Multi-Light System: Sun/Moon, sky/ground ambient, 8 point lights (braziers)",
        "Soft Shadow Mapping: 2048² Depth FBO pass with 3×3 PCF & depth bias ('X')",
        "Real-Time Ray Tracing: Kay-Kajiya slab intersection test on GPU ('K')",
        "Full Compound Fortress: 4 towers, double curtain walls, river, drawbridge",
        "Autonomous Tactical Siege: 6-phase battle, archers, shields, destruction"
    ]
    for item in items_delivered:
        pi = tf_d.add_paragraph()
        pi.text = "✓ " + item
        pi.font.size = Pt(11)
        pi.font.color.rgb = TEXT_LIGHT
        pi.space_before = Pt(6)

    # =========================================================================
    # SLIDE 3: System Architecture & Graphics Pipeline
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "System Architecture & Real-Time Graphics Pipeline")

    phases = [
        ("1. Input & Camera", "Poll GLFW keyboard & mouse\nUpdate 360° spherical orbit camera\nCalculate deltaTime & frame pacing", ACCENT_BLUE),
        ("2. Physical Simulation", "Semi-implicit Euler ballistics\nSphere-vs-AABB continuous collision\nAutonomous 6-phase state machine", ACCENT_GOLD),
        ("3. Shadow Depth Pass", "Render light-space depth map\n2048×2048 32-bit floating-point FBO\nCull front faces to prevent acne", ACCENT_RED),
        ("4. Main Render Pass", "Activate active shader (Gouraud/Phong)\nTraverse hierarchical scene graph\nUpload lighting, PCF shadows & ray boxes", ACCENT_GREEN)
    ]

    for idx, (title, desc, col) in enumerate(phases):
        x = Inches(0.8 + idx * 2.95)
        add_card(s3, x, Inches(1.8), Inches(2.8), Inches(3.2))
        tb = s3.shapes.add_textbox(x + Inches(0.15), Inches(1.9), Inches(2.5), Inches(3.0))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_LIGHT
        p2.space_before = Pt(8)

    # Bottom summary card
    add_card(s3, Inches(0.8), Inches(5.3), Inches(11.7), Inches(1.7))
    tb_arch = s3.shapes.add_textbox(Inches(1.0), Inches(5.4), Inches(11.3), Inches(1.5))
    tf_arch = tb_arch.text_frame
    tf_arch.word_wrap = True
    p = tf_arch.paragraphs[0]
    p.text = "Modular C++ Architecture (Over 40 Modular Classes):"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GOLD

    p2 = tf_arch.add_paragraph()
    p2.text = "• Core Math & Pipeline: Transform, Mesh (VAO/VBO/EBO), shaderClass, Primitives, Dimensions, Palette\n• Scene Graph Entities: Cannon (Carriage, Wheel, Shaft), Projectile, CompoundCastle, Tower, Door, Bridge, Water\n• Characters & Dynamic Systems: Soldier, Archer, Arrow, ShadowMap, ParticleSystem, Scenery, SkyClouds, Birds"
    p2.font.size = Pt(10)
    p2.font.color.rgb = TEXT_LIGHT
    p2.space_before = Pt(4)

    # =========================================================================
    # SLIDE 4: Hierarchical Transformations & Kinematics
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Hierarchical Transformations & Kinematics")

    # Math Card
    add_card(s4, Inches(0.8), Inches(1.8), Inches(6.0), Inches(5.2))
    tb_m = s4.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(5.6), Inches(5.0))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True
    p = tf_m.paragraphs[0]
    p.text = "📐 KINEMATIC FORMULATION & EQUATIONS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GOLD

    eqs = [
        ("Affine Model Matrix", "M_model = T(t) • R_y(ψ) • R_z(θ) • S(s)"),
        ("Non-Slip Rolling Constraint", "Δθ_wheel = −Δs / R_wheel  (R = 0.584 m)\nGuarantees zero linear tire skid across ground"),
        ("Trunnion Elevation Pivot", "M_barrel = T(p_pivot) • R_z(θ_elev) • T(−p_pivot) • T(x_recoil)\nRotates about arbitrary hinge; clamped 0° to 45°"),
        ("Muzzle Spawn Projection", "p_muzzle = M_carriage • M_barrel • [L_barrel, 0, 0, 1]^T\nExact world-space spawn coordinate for ballistics"),
        ("View & Perspective Matrices", "V = lookAt(E, C, Up),  P = perspective(45°, aspect, 0.1, 600)")
    ]
    for title, formula in eqs:
        pt = tf_m.add_paragraph()
        pt.text = "• " + title + ":"
        pt.font.size = Pt(11)
        pt.font.bold = True
        pt.font.color.rgb = ACCENT_BLUE
        pt.space_before = Pt(5)

        pf = tf_m.add_paragraph()
        pf.text = "   " + formula
        pf.font.size = Pt(10)
        pf.font.color.rgb = TEXT_LIGHT

    # Right Card: Image plates
    add_card(s4, Inches(7.1), Inches(1.8), Inches(5.4), Inches(5.2))
    tb_i = s4.shapes.add_textbox(Inches(7.3), Inches(1.9), Inches(5.0), Inches(0.4))
    tb_i.text_frame.paragraphs[0].text = "ELEVATION STAGES (0°, 25°, 45°)"
    tb_i.text_frame.paragraphs[0].font.size = Pt(12)
    tb_i.text_frame.paragraphs[0].font.bold = True
    tb_i.text_frame.paragraphs[0].font.color.rgb = ACCENT_GOLD

    img_elev0 = os.path.join("docs", "images", "05_elev_00.png")
    img_elev25 = os.path.join("docs", "images", "06_elev_25.png")
    img_elev45 = os.path.join("docs", "images", "07_elev_45.png")
    if os.path.exists(img_elev0):
        s4.shapes.add_picture(img_elev0, Inches(7.3), Inches(2.4), width=Inches(5.0))

    # =========================================================================
    # SLIDE 5: Shading Models: Gouraud vs. Blinn-Phong
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Dual Shading Models: Gouraud vs. Blinn-Phong (Key 'G')")

    # Left: Gouraud
    add_card(s5, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.2))
    tb_g = s5.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(5.3), Inches(5.0))
    tf_g = tb_g.text_frame
    tf_g.word_wrap = True
    p = tf_g.paragraphs[0]
    p.text = "GOURAUD SHADING (PER-VERTEX)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GOLD

    g_points = [
        "Implementation: gouraud.vert + gouraud.frag",
        "Evaluation: Lighting calculated strictly at vertices in vertex shader",
        "Interpolation: Hardware rasterizer linearly interpolates RGB colors",
        "Equation: I_v = I_amb + I_diff(N•L) + I_spec(R•V)^α",
        "Characteristics: Ultra-fast execution, O(V) ALU cost",
        "Artifacts: Prone to Mach banding; misses specular highlights if they land inside large polygons"
    ]
    for pt in g_points:
        pi = tf_g.add_paragraph()
        pi.text = "• " + pt
        pi.font.size = Pt(10)
        pi.font.color.rgb = TEXT_LIGHT
        pi.space_before = Pt(4)

    img_gouraud = os.path.join("docs", "images", "14_shading_gouraud.png")
    if os.path.exists(img_gouraud):
        s5.shapes.add_picture(img_gouraud, Inches(1.0), Inches(4.7), width=Inches(5.3))

    # Right: Phong / Blinn-Phong
    add_card(s5, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.2))
    tb_p = s5.shapes.add_textbox(Inches(7.0), Inches(1.9), Inches(5.3), Inches(5.0))
    tf_p = tb_p.text_frame
    tf_p.word_wrap = True
    p = tf_p.paragraphs[0]
    p.text = "BLINN-PHONG SHADING (PER-FRAGMENT)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    p_points = [
        "Implementation: lit.vert + lit.frag (ACTIVE DEFAULT)",
        "Evaluation: Normal & world position interpolated; full lighting evaluated per pixel in fragment shader",
        "Halfway Vector: H = (L + V) / ||L + V||",
        "Equation: I_f = I_amb + (N•L)•C_diff + (N•H)^α • C_spec + I_rim",
        "Fresnel Term: I_rim = k_metal • (1 − max(N•V, 0))^3.5 • C_rim",
        "Characteristics: Smooth elliptical highlights, eliminates Mach banding, supports procedural stone/wood"
    ]
    for pt in p_points:
        pi = tf_p.add_paragraph()
        pi.text = "• " + pt
        pi.font.size = Pt(10)
        pi.font.color.rgb = TEXT_LIGHT
        pi.space_before = Pt(4)

    img_phong = os.path.join("docs", "images", "13_shading_phong.png")
    if os.path.exists(img_phong):
        s5.shapes.add_picture(img_phong, Inches(7.0), Inches(4.7), width=Inches(5.3))

    # =========================================================================
    # SLIDE 6: Multi-Light Illumination & Atmospheric Effects
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Multi-Light System & Atmospheric Rendering")

    light_features = [
        ("Hemispheric Ambient", "I_amb = mix(C_ground, C_sky, (N_y + 1)/2) • k_a\nWarm ground reflections + cool celestial skylight.", ACCENT_BLUE),
        ("Directional Sun & Moon", "Dynamic sunlight (day) & pale moonlight (night).\nToggle in real time with key 'N'. Includes visible emissive celestial bodies.", ACCENT_GOLD),
        ("8 Quadratic Point Lights", "f_att(d) = 1 / (1 + 0.12d + 0.045d²)\nApplied to castle wall braziers, torches, and explosive cannon muzzle sparks.", ACCENT_RED),
        ("Distance Fog & ACES", "Exponential fog: 1 − exp(−(d_fog • k)^1.35)\nACES Filmic Tone Mapping + Gamma 2.2 correction.", ACCENT_GREEN)
    ]
    for idx, (title, desc, col) in enumerate(light_features):
        x = Inches(0.8 + idx * 2.95)
        add_card(s6, x, Inches(1.8), Inches(2.8), Inches(5.2))
        tb = s6.shapes.add_textbox(x + Inches(0.15), Inches(1.9), Inches(2.5), Inches(5.0))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_LIGHT
        p2.space_before = Pt(8)

    # =========================================================================
    # SLIDE 7: Soft Shadow Mapping (PCF) & Depth Bias
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Soft Shadow Mapping with 3x3 PCF (Key 'X')")

    # Math Card
    add_card(s7, Inches(0.8), Inches(1.8), Inches(6.0), Inches(5.2))
    tb_s = s7.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(5.6), Inches(5.0))
    tf_s = tb_s.text_frame
    tf_s.word_wrap = True
    p = tf_s.paragraphs[0]
    p.text = "🎯 SHADOW PIPELINE & PCF ALGORITHM"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GOLD

    s_items = [
        "Pass 1 (Light Space): Render depth to 2048² 32-bit floating-point FBO texture",
        "Light Projection: Orthographic frustum covering entire battlefield",
        "Slope-Scaled Depth Bias: Eliminates shadow acne on grazing surfaces:\nbias = max(0.0035 • (1.0 − N•L), 0.0006)",
        "Percentage-Closer Filtering (PCF): 3×3 kernel averaging 9 depth samples:\nshadow = (1/9) Σ Σ (proj_z − bias > depth ? 0.0 : 1.0)",
        "Result: Visually realistic soft penumbrae under wheels, fortress walls, and soldiers"
    ]
    for item in s_items:
        pi = tf_s.add_paragraph()
        pi.text = "• " + item
        pi.font.size = Pt(11)
        pi.font.color.rgb = TEXT_LIGHT
        pi.space_before = Pt(6)

    # Screenshot Card
    add_card(s7, Inches(7.1), Inches(1.8), Inches(5.4), Inches(5.2))
    img_shadow = os.path.join("docs", "images", "patrol_sidewalls.png")
    if os.path.exists(img_shadow):
        s7.shapes.add_picture(img_shadow, Inches(7.3), Inches(2.2), width=Inches(5.0))

    # =========================================================================
    # SLIDE 8: Real-Time Ray Tracing (Bonus Marks Feature)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Real-Time Ray Tracing: Kay-Kajiya Slab Method (Key 'K')")

    add_card(s8, Inches(0.8), Inches(1.8), Inches(6.0), Inches(5.2))
    tb_rt = s8.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(5.6), Inches(5.0))
    tf_rt = tb_rt.text_frame
    tf_rt.word_wrap = True
    p = tf_rt.paragraphs[0]
    p.text = "⚡ GPU SLAB RAY-BOX INTERSECTION"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GOLD

    rt_points = [
        "Target Feature: Real-time ray tracing implemented directly on GPU in GLSL",
        "Ray Formulation: r(t) = o + t•d, where o = p_frag + 0.05•N, d = lightDir",
        "Kay-Kajiya Slab Algorithm: Evaluates ray intersection across 3 pairs of parallel AABB planes:\nt0 = (bMin − ro) / rd,  t1 = (bMax − ro) / rd\ntNear = max(tmin.x, tmin.y, tmin.z),  tFar = min(tmax.x, tmax.y, tmax.z)",
        "Hit Condition: (tNear ≤ tFar) && (tFar > 0.0)",
        "Obstacle Array: Active castle walls & towers uploaded as uniform array of 36 boxes",
        "Key 'K' Toggle: Real-time switch casting exact, razor-sharp ray-traced shadows without texture discretization!"
    ]
    for item in rt_points:
        pi = tf_rt.add_paragraph()
        pi.text = "• " + item
        pi.font.size = Pt(11)
        pi.font.color.rgb = TEXT_LIGHT
        pi.space_before = Pt(5)

    # GLSL Snippet Card
    add_card(s8, Inches(7.1), Inches(1.8), Inches(5.4), Inches(5.2))
    tb_code = s8.shapes.add_textbox(Inches(7.3), Inches(1.9), Inches(5.0), Inches(5.0))
    tf_code = tb_code.text_frame
    tf_code.word_wrap = True
    p = tf_code.paragraphs[0]
    p.text = "GLSL SHADOW RAY KERNEL"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    code_text = (
        "// In lit.frag:\n"
        "float calcRayShadow(vec3 ro, vec3 rd, float maxDist) {\n"
        "    for (int i = 0; i < numBoxes; i++) {\n"
        "        float tHit;\n"
        "        if (hitBox(ro, rd, boxMin[i], boxMax[i], tHit)) {\n"
        "            if (tHit > 0.15 && tHit < maxDist) {\n"
        "                return 0.0; // Occluded!\n"
        "            }\n"
        "        }\n"
        "    }\n"
        "    return 1.0; // Fully visible\n"
        "}\n\n"
        "// Live runtime toggle:\n"
        "// Press 'K' to toggle Ray Tracing mode\n"
        "// Press 'X' to return to Soft PCF Depth Map"
    )
    pi = tf_code.add_paragraph()
    pi.text = code_text
    pi.font.size = Pt(10)
    pi.font.color.rgb = ACCENT_GOLD
    pi.font.name = "Courier New"
    pi.space_before = Pt(6)

    # =========================================================================
    # SLIDE 9: Physics Simulation, Ballistics & Dynamic Destruction
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Physics Simulation, Ballistics & Progressive Destruction")

    # Math Card
    add_card(s9, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.2))
    tb_p = s9.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(5.3), Inches(5.0))
    tf_p = tb_p.text_frame
    tf_p.word_wrap = True
    p = tf_p.paragraphs[0]
    p.text = "💥 BALLISTIC PHYSICS & COLLISION"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GOLD

    p_items = [
        ("Semi-Implicit Euler Integration", "v(t + Δt) = v(t) + g • Δt\np(t + Δt) = p(t) + v(t + Δt) • Δt\nEnergy-conserving symplectic numerical integrator"),
        ("Sphere-vs-AABB Collision", "c_i = clamp(p_i, bMin_i, bMax_i)\nHit if ||p_sphere − c|| ≤ r_sphere"),
        ("Wall Fracture & Chipping", "10% health per impact; color darkens on hit;\nboundaries removed at 0 HP so balls fly through gap"),
        ("Door Splintering", "Wooden gates shatter completely after 2 cannon hits"),
        ("Drawbridge Physics", "Bridge chains snap when shot; timber deck falls 85° → 0°")
    ]
    for title, desc in p_items:
        pt = tf_p.add_paragraph()
        pt.text = "• " + title + ":"
        pt.font.size = Pt(11)
        pt.font.bold = True
        pt.font.color.rgb = ACCENT_BLUE
        pt.space_before = Pt(4)

        pd = tf_p.add_paragraph()
        pd.text = "   " + desc
        pd.font.size = Pt(10)
        pd.font.color.rgb = TEXT_LIGHT

    # Image Card
    add_card(s9, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.2))
    img_breach = os.path.join("docs", "images", "cannonsiege_breach.png")
    if os.path.exists(img_breach):
        s9.shapes.add_picture(img_breach, Inches(7.0), Inches(2.2), width=Inches(5.3))

    # =========================================================================
    # SLIDE 10: Autonomous Battlefield Simulation & Character Animation
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Autonomous 6-Phase Tactical Siege State Machine (Key 'B')")

    phases_sim = [
        ("Phase 0: Peaceful Patrol", "Sentries pace side walls.\nRelief guards climb wooden ladders periodically.", ACCENT_BLUE),
        ("Phase 1: Advance & Alarm", "Alarm rings; defenders raise drawbridge (0°→85°).\nCannons & infantry march in synchronized lockstep.", ACCENT_GOLD),
        ("Phase 2: Archer Barrage", "Castle archers loose arrows; infantry raise shields.\nCannoneer 1 hit while priming; relieved by infantryman.", ACCENT_RED),
        ("Phase 3: Severing Hinge", "Cannoneer 2 (preserved unharmed) shoots bridge chains.\nChains snap; drawbridge slams flat across moat.", ACCENT_BLUE),
        ("Phase 4: Coordinated Siege", "Simultaneous 3-cannon barrage:\nCannon 1 hits left wall, Cannon 3 hits right, Cannon 2 breaks door.", ACCENT_GOLD),
        ("Phase 5: Courtyard Charge", "Infantry form 2-column bridge crossing;\ncharge courtyard; animated melee duel; golden crest victory!", ACCENT_GREEN)
    ]
    for idx, (title, desc, col) in enumerate(phases_sim):
        row = idx // 3
        col_idx = idx % 3
        x = Inches(0.8 + col_idx * 3.95)
        y = Inches(1.8 + row * 2.65)
        add_card(s10, x, y, Inches(3.8), Inches(2.45))
        tb = s10.shapes.add_textbox(x + Inches(0.15), y + Inches(0.15), Inches(3.5), Inches(2.1))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_LIGHT
        p2.space_before = Pt(6)

    # =========================================================================
    # SLIDE 11: User Interaction & Comprehensive Control Scheme
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "User Interaction & Comprehensive Controls")

    add_card(s11, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.2))
    tb_c = s11.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(11.3), Inches(5.0))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True

    p = tf_c.paragraphs[0]
    p.text = "INTERACTIVE KEYBOARD & MOUSE CONTROLS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GOLD

    controls = [
        ("Mouse Left Drag", "360° Spherical orbit camera rotation (yaw & pitch)"),
        ("Mouse Wheel", "Interactive zoom in / zoom out (radius 12m to 240m)"),
        ("Left / Right Arrows", "Drive active cannon forward / backward (wheels roll with no slip)"),
        ("Up / Down Arrows", "Elevate / depress active cannon barrel (clamped 0° to 45°)"),
        ("W / S Keys", "Steer / yaw active cannon barrel left / right"),
        ("Spacebar", "Fire active cannon (triggers recoil, ballistic ball, muzzle sparks)"),
        ("Keys 1, 2, 3 / A", "Select Cannon 1 (left), Cannon 2 (center), Cannon 3 (right), or All (A)"),
        ("G Key", "Toggle Shading Mode in real time: Gouraud (per-vertex) <-> Blinn-Phong"),
        ("X / Y Keys", "Toggle Soft Shadow Mapping: Enable/disable 3×3 PCF depth FBO shadows"),
        ("K Key", "Toggle Real-Time Ray Tracing: Enable GPU ray-box slab shadow rays"),
        ("R Key (Hold)", "Manually raise drawbridge (0° → 85°); release to lower"),
        ("N Key", "Toggle Day / Night: Switch solar/lunar illumination and sky"),
        ("B / P / T Keys", "Start battle simulation (B), Pause/resume (P), Reset simulation (T)")
    ]

    for key, desc in controls:
        pi = tf_c.add_paragraph()
        pi.text = f"• {key:20s} : {desc}"
        pi.font.size = Pt(10)
        pi.font.color.rgb = TEXT_LIGHT
        pi.space_before = Pt(2)

    # =========================================================================
    # SLIDE 12: Project Demo Video Breakdown (presentation.mp4)
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "Project Demo Video Breakdown (presentation.mp4)")

    timeline = [
        ("0:00 - 0:05", "Peaceful Patrol & Wall Sentries", "Sentries patrol curtain walls; soft PCF shadows visible across grass; day sky."),
        ("0:05 - 0:10", "Alarm & Synchronized Advance", "Defenders winch up drawbridge; 3 cannons & infantry battalion march in formation."),
        ("0:10 - 0:18", "Archer Barrage & Relief", "Archers loose volleys; troops raise shields; cannoneer 1 hit & replaced by infantry."),
        ("0:18 - 0:25", "Severing the Bridge Hinge", "Center cannon fires precision shot; drawbridge chains snap; bridge drops across moat."),
        ("0:25 - 0:35", "Coordinated Artillery Bombardment", "Cannons 1 & 3 pound curtain walls; Cannon 2 shatters wooden gatehouse doors."),
        ("0:35 - 0:48", "Courtyard Charge & Melee Victory", "Infantry form 2-column bridge march, breach courtyard, win melee, gold crest pulses!")
    ]

    for idx, (timeframe, title, details) in enumerate(timeline):
        x = Inches(0.8 + (idx % 2) * 5.95)
        y = Inches(1.8 + (idx // 2) * 1.75)
        add_card(s12, x, y, Inches(5.7), Inches(1.6))
        tb = s12.shapes.add_textbox(x + Inches(0.15), y + Inches(0.1), Inches(5.4), Inches(1.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"⏱ {timeframe} — {title}"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = ACCENT_GOLD

        p2 = tf.add_paragraph()
        p2.text = details
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_LIGHT
        p2.space_before = Pt(4)

    # =========================================================================
    # SLIDE 13: High-Resolution Screenshot Showcase
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_background(s13)
    add_header(s13, "Visual Results & Screenshot Gallery")

    shots = [
        ("Full Scene Overview & March", "alarm_march.png"),
        ("Archer Barrage & Shield Wall", "cannoneer1_hit.png"),
        ("Artillery Breach of Gatehouse", "cannonsiege_breach.png"),
        ("Courtyard Breach & Melee", "charge_troops_inside.png")
    ]

    for idx, (caption, fname) in enumerate(shots):
        x = Inches(0.8 + (idx % 2) * 5.95)
        y = Inches(1.8 + (idx // 2) * 2.65)
        add_card(s13, x, y, Inches(5.7), Inches(2.5))
        img_path = os.path.join("docs", "images", fname)
        if os.path.exists(img_path):
            s13.shapes.add_picture(img_path, x + Inches(0.15), y + Inches(0.15), width=Inches(3.2))

        tb = s13.shapes.add_textbox(x + Inches(3.45), y + Inches(0.3), Inches(2.1), Inches(1.9))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = caption
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ACCENT_GOLD

    # =========================================================================
    # SLIDE 14: Conclusion & Academic Learnings
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    set_slide_background(s14)
    add_header(s14, "Conclusion & Academic Takeaways")

    add_card(s14, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.2))
    tb_concl = s14.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(11.3), Inches(5.0))
    tf_concl = tb_concl.text_frame
    tf_concl.word_wrap = True

    p = tf_concl.paragraphs[0]
    p.text = "KEY ACADEMIC ACHIEVEMENTS & LEARNINGS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GOLD

    concl_points = [
        "100% Scope Fulfillment: Delivered every baseline requirement and expanded into an advanced 3-D siege simulator.",
        "Mathematical Rigor: Exact affine transformations, forward kinematic scene graph, and non-slipping rolling constraints.",
        "Shading Mastery: Deep comparative analysis and real-time switching between Gouraud and Blinn-Phong shading.",
        "Advanced Illumination: Multi-light system, quadratic distance attenuation, and hemispherical ambient bounce.",
        "Dual Shadow Implementations: Real-time 3×3 PCF soft shadow mapping and GPU ray-box slab ray tracing.",
        "Physical Simulation: Symplectic semi-implicit Euler integration, continuous AABB collisions, and multi-stage fracture.",
        "Dynamic Interactivity: Full orbit camera control, manual artillery firing, and interactive drawbridge winches.",
        "Academic References Followed: Joey de Vries (LearnOpenGL), Blinn (1977), Phong (1975), Kay & Kajiya (1986)."
    ]
    for pt in concl_points:
        pi = tf_concl.add_paragraph()
        pi.text = "✓ " + pt
        pi.font.size = Pt(11)
        pi.font.color.rgb = TEXT_LIGHT
        pi.space_before = Pt(5)

    # =========================================================================
    # SLIDE 15: Thank You Slide
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    set_slide_background(s15)

    add_card(s15, Inches(1.5), Inches(1.8), Inches(10.333), Inches(4.2))
    tb_ty = s15.shapes.add_textbox(Inches(1.8), Inches(2.2), Inches(9.7), Inches(3.4))
    tf_ty = tb_ty.text_frame
    tf_ty.word_wrap = True

    p = tf_ty.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "THANK YOU!"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GOLD

    p2 = tf_ty.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "Medieval Fortress Siege Simulation"
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_before = Pt(10)

    p3 = tf_ty.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    p3.text = "Computer Graphics Lab Project Showcase"
    p3.font.size = Pt(15)
    p3.font.color.rgb = ACCENT_BLUE
    p3.space_before = Pt(6)

    p4 = tf_ty.add_paragraph()
    p4.alignment = PP_ALIGN.CENTER
    p4.text = "Ready for Live Demo, Parameter Adjustments & Questions"
    p4.font.size = Pt(13)
    p4.font.color.rgb = TEXT_MUTED
    p4.space_before = Pt(16)

    # Save to presentation directory
    os.makedirs("presentation", exist_ok=True)
    out_path1 = os.path.join("presentation", "presentation.pptx")
    out_path2 = os.path.join("presentation", "Medieval_Battle_Showcase.pptx")
    prs.save(out_path1)
    prs.save(out_path2)
    print(f"Presentation successfully saved to: {out_path1} and {out_path2}")

if __name__ == "__main__":
    build_presentation()
